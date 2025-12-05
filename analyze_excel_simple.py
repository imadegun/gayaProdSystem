
import openpyxl
import sys
import os

file_path = "docs/directory_list.xlsx"

if not os.path.exists(file_path):
    print(f"File not found: {file_path}")
    sys.exit(1)

try:
    wb = openpyxl.load_workbook(file_path, read_only=True)
    print("Sheets found:", wb.sheetnames)
    
    for sheet_name in wb.sheetnames:
        print(f"\n--- Sheet: {sheet_name} ---")
        ws = wb[sheet_name]
        
        # Get headers (first row)
        headers = []
        for cell in ws[1]:
            headers.append(cell.value)
        print("Headers:", headers)
        
        # Get first row of data (second row)
        first_row_data = []
        for cell in ws[2]:
            first_row_data.append(cell.value)
        print("First row data:", first_row_data)
            
except Exception as e:
    print(f"Error reading excel file: {e}")
